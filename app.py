# -*- coding: utf-8 -*-
"""TerreBalance : atelier pedagogique de volumes et de transport des terres."""
from __future__ import annotations

import copy
import hashlib
import io
import json
import math
from datetime import datetime
from html import escape
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import plotly.graph_objects as go
import streamlit as st
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Image, KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from scipy import stats


GREEN, BLUE, CORAL = "#166534", "#2563EB", "#D94841"
INK, MUTED, BG = "#202427", "#56616B", "#F5F7FA"
PIT_KEYS = ("longueur_m", "largeur_m", "profondeur_m", "foisonnement")
ROW_KEYS = ("identifiant",) + PIT_KEYS
ORIGINS = ("Saisie utilisateur", "Exemple fictif", "Exemple fictif modifié", "Import : origine à documenter")
LIMITS = "Modèle pédagogique : parois verticales, sans vérification de stabilité, portance, charge en masse, contraintes routières, compactage ou tassement."
SECTIONS = ["Terrassement", "Plusieurs fouilles", "Comparaison", "Incertitudes", "Rapport"]


def number(value, name, low, high, integer=False):
    if isinstance(value, (bool, np.bool_)) or not isinstance(value, (int, float, np.integer, np.floating)):
        raise ValueError(f"{name} : un nombre est requis.")
    if not math.isfinite(float(value)) or not low <= value <= high:
        raise ValueError(f"{name} : valeur attendue entre {low:g} et {high:g}.")
    if integer and int(value) != value:
        raise ValueError(f"{name} : un entier est requis.")
    return int(value) if integer else float(value)


def label(value, name, maximum=80):
    if not isinstance(value, str) or not value.strip() or len(value.strip()) > maximum:
        raise ValueError(f"{name} : texte non vide limité à {maximum} caractères.")
    if any(ord(char) < 32 for char in value):
        raise ValueError(f"{name} : caractères de contrôle interdits.")
    return value.strip()


def validate_pit(pit):
    if not isinstance(pit, dict) or set(pit) != set(PIT_KEYS):
        raise ValueError("Fouille : quatre paramètres géométriques sont requis.")
    return dict(
        longueur_m=number(pit["longueur_m"], "Longueur (m)", .01, 1000),
        largeur_m=number(pit["largeur_m"], "Largeur (m)", .01, 1000),
        profondeur_m=number(pit["profondeur_m"], "Profondeur (m)", 0, 100),
        foisonnement=number(pit["foisonnement"], "Coefficient k", 1, 3),
    )


def default_pit():
    return dict(longueur_m=10., largeur_m=5., profondeur_m=1., foisonnement=1.2)


def new_project():
    return dict(application="TerreBalance", version=1, titre="Mon étude de terrassement",
        origine="Saisie utilisateur", fouille=None, capacite_m3=10., fouilles=[],
        origine_fouilles="Saisie utilisateur", capacite_lot_m3=10.,
        variante=dict(profondeur_m=1.2, foisonnement=1.2, capacite_m3=10.), comparaison_active=False,
        incertitude=dict(min_m=.9, mode_m=1., max_m=1.1, tirages=10000, graine=42, limite_voyages=6), simulation_active=False)


def validate_rows(rows):
    if not isinstance(rows, list) or len(rows) > 100:
        raise ValueError("Le lot doit être une liste de 100 fouilles au maximum.")
    output, identifiers = [], set()
    for i, row in enumerate(rows, 1):
        if not isinstance(row, dict) or set(row) != set(ROW_KEYS):
            raise ValueError(f"Ligne {i} : colonnes attendues {', '.join(ROW_KEYS)}.")
        ident = label(row["identifiant"], f"Ligne {i}, identifiant", 40)
        if ident in identifiers:
            raise ValueError(f"Identifiant dupliqué : {ident}.")
        identifiers.add(ident)
        output.append(dict(identifiant=ident, **validate_pit({k:row[k] for k in PIT_KEYS})))
    return output


def validate_project(project):
    if not isinstance(project, dict) or set(project) != set(new_project()):
        raise ValueError("Structure JSON de projet incomplète ou inconnue.")
    c = copy.deepcopy(project)
    if c["application"] != "TerreBalance" or type(c["version"]) is not int or c["version"] != 1:
        raise ValueError("Projet TerreBalance version 1 attendu.")
    c["titre"] = label(c["titre"], "Nom du projet")
    for key in ("origine", "origine_fouilles"):
        if c[key] not in ORIGINS:
            raise ValueError("Origine des données inconnue.")
    if c["fouille"] is not None:
        c["fouille"] = validate_pit(c["fouille"])
    c["fouilles"] = validate_rows(c["fouilles"])
    for key in ("capacite_m3", "capacite_lot_m3"):
        c[key] = number(c[key], "Capacité utile (m³)", .1, 1000)
    v = c["variante"]
    if not isinstance(v, dict) or set(v) != {"profondeur_m", "foisonnement", "capacite_m3"}:
        raise ValueError("Paramètres de variante invalides.")
    c["variante"] = dict(profondeur_m=number(v["profondeur_m"], "Profondeur variante", 0, 100),
        foisonnement=number(v["foisonnement"], "Coefficient variante", 1, 3),
        capacite_m3=number(v["capacite_m3"], "Capacité variante", .1, 1000))
    u = c["incertitude"]
    if not isinstance(u, dict) or set(u) != {"min_m", "mode_m", "max_m", "tirages", "graine", "limite_voyages"}:
        raise ValueError("Paramètres de simulation invalides.")
    u = {k:number(u[k], k, 0, 100) for k in ("min_m", "mode_m", "max_m")} | {
        "tirages":number(u["tirages"], "Tirages", 100, 50000, True),
        "graine":number(u["graine"], "Graine", 0, 2147483647, True),
        "limite_voyages":number(u["limite_voyages"], "Limite de voyages", 0, 3000000000, True)}
    if not u["min_m"] <= u["mode_m"] <= u["max_m"]:
        raise ValueError("Profondeurs : minimum ≤ mode ≤ maximum requis.")
    c["incertitude"] = u
    for key in ("simulation_active", "comparaison_active"):
        if type(c[key]) is not bool:
            raise ValueError("Les états de calcul doivent être booléens.")
        if c[key] and c["fouille"] is None:
            raise ValueError("Une fouille enregistrée est nécessaire avant une comparaison ou une simulation.")
    return c


def payload(project):
    return json.dumps(project, sort_keys=True, ensure_ascii=False, allow_nan=False)


def stable_ratio(volume, capacity):
    # Snap only floating-point noise (8 ULP), never a genuinely positive volume to zero.
    ratio = np.asarray(volume, dtype=float) / capacity
    rounded = np.rint(ratio)
    tolerance = 8 * np.spacing(np.maximum(np.abs(ratio), 1.))
    return np.where((rounded > 0) & (np.abs(ratio-rounded) <= tolerance), rounded, ratio)


def trip_counts(volume, capacity):
    capacity = number(capacity, "Capacité utile (m³)", .1, 1000)
    values = np.asarray(volume, dtype=float)
    if not np.all(np.isfinite(values)) or np.any(values < 0) or np.any(values > 3e10):
        raise ValueError("Volume à transporter invalide.")
    return np.ceil(stable_ratio(values, capacity)).astype(np.int64)


def transport(volume, capacity):
    n = int(trip_counts(volume, capacity))
    last = float(np.clip(volume-(n-1)*capacity, 0, capacity)) if n else 0.
    if n and stable_ratio(volume, capacity) == n:
        last = float(capacity)
    partial = n > 0 and last < capacity
    return dict(voyages=n, dernier_m3=last, dernier_pct=100*last/capacity if n else 0.,
                complets=n-int(partial), partiel=partial)


def calculate(pit, capacity):
    p = validate_pit(pit)
    capacity = number(capacity, "Capacité utile (m³)", .1, 1000)
    raw = p["longueur_m"]*p["largeur_m"]*p["profondeur_m"]
    loose = raw*p["foisonnement"]
    return dict(en_place_m3=raw, foisonne_m3=loose, augmentation_m3=loose-raw,
                augmentation_pct=100*(p["foisonnement"]-1), capacite_m3=capacity, **transport(loose, capacity))


def batch_results(rows, capacity):
    rows = validate_rows(rows)
    capacity = number(capacity, "Capacité du lot (m³)", .1, 1000)
    records = [dict(**row, **calculate({k:row[k] for k in PIT_KEYS}, capacity)) for row in rows]
    loose = math.fsum(r["foisonne_m3"] for r in records)
    return dict(table=pd.DataFrame(records), en_place_m3=math.fsum(r["en_place_m3"] for r in records),
        foisonne_m3=loose, separes=sum(r["voyages"] for r in records), **transport(loose, capacity))


def compare(project):
    p = project["fouille"]
    if p is None:
        raise ValueError("Enregistrer la fouille principale.")
    v = project["variante"]
    other = p | {"profondeur_m":v["profondeur_m"], "foisonnement":v["foisonnement"]}
    return calculate(p, project["capacite_m3"]), calculate(other, v["capacite_m3"])


def exact_interval(k, n):
    return (0. if k == 0 else float(stats.beta.ppf(.025, k, n-k+1)),
            1. if k == n else float(stats.beta.ppf(.975, k+1, n-k)))


def simulate(project):
    c = validate_project(project)
    if c["fouille"] is None:
        raise ValueError("Enregistrer la fouille avant de simuler.")
    p, u = c["fouille"], c["incertitude"]
    a, m, b = u["min_m"], u["mode_m"], u["max_m"]
    factor = p["longueur_m"]*p["largeur_m"]*p["foisonnement"]
    capacity, n, limit = c["capacite_m3"], u["tirages"], u["limite_voyages"]
    if a == b:
        depths = np.full(n, a)
        analytic = float(int(trip_counts(a*factor, capacity)) > limit)
        quantiles = np.full(3, a*factor)
    else:
        law = stats.triang((m-a)/(b-a), loc=a, scale=b-a)
        depths = law.rvs(size=n, random_state=np.random.default_rng(u["graine"]))
        # The analytic threshold follows the same ULP convention as integer transport.
        threshold = float(limit)+8*np.spacing(float(max(limit, 1))) if limit else 0.
        analytic = float(law.sf(threshold*capacity/factor))
        quantiles = law.ppf([.05, .5, .95])*factor
    volumes = depths*factor
    trips = trip_counts(volumes, capacity)
    exceed = trips > limit
    if not np.array_equal(exceed, stable_ratio(volumes, capacity) > limit):
        raise ArithmeticError("Contrôle voyages/volume incohérent.")
    low, high = exact_interval(int(exceed.sum()), n)
    return dict(table=pd.DataFrame({"tirage":np.arange(1,n+1), "profondeur_m":depths,
        "volume_foisonne_m3":volumes, "voyages":trips, "depasse_limite":exceed}),
        probabilite=float(exceed.mean()), analytique=analytic, ic_bas=low, ic_haut=high,
        moyenne_volume=float(volumes.mean()), moyenne_voyages=float(trips.mean()),
        volume_moyen_theorique=factor*(a+m+b)/3, quantiles=np.quantile(volumes,[.05,.5,.95]),
        quantiles_theoriques=quantiles, constant=a == b)


@st.cache_data(max_entries=6, ttl=1800, show_spinner=False)
def cached_simulation(configuration):
    return simulate(json.loads(configuration))


def example_project():
    c = new_project()
    c.update(titre="Fouilles de démonstration", origine="Exemple fictif", fouille=default_pit(), origine_fouilles="Exemple fictif")
    c["fouilles"] = [dict(identifiant=f"F{i}", longueur_m=5., largeur_m=2., profondeur_m=1., foisonnement=1.2) for i in (1,2)]
    return validate_project(c)


def safe_csv(frame):
    df = frame.copy()
    for col in df.columns:
        if pd.api.types.is_object_dtype(df[col].dtype) or pd.api.types.is_string_dtype(df[col].dtype):
            df[col] = df[col].map(lambda x:"'"+x if isinstance(x,str) and x.startswith(("=","+","-","@")) else x)
    return df.to_csv(index=False).encode("utf-8-sig")


@st.cache_data(max_entries=8,ttl=1800,show_spinner=False)
def excel_bytes(sheets):
    wb = Workbook()
    wb.remove(wb.active)
    for name, frame in sheets.items():
        ws = wb.create_sheet(name)
        ws.append([str(c) for c in frame.columns])
        for row in frame.itertuples(index=False, name=None):
            ws.append([v.item() if isinstance(v,np.generic) else v for v in row])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for row in ws:
            for cell in row:
                if isinstance(cell.value,str):
                    cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="166534")
            cell.font = Font(color="FFFFFF", bold=True)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(36,max(16,len(str(col[0].value))+2))
    out = io.BytesIO()
    wb.save(out)
    wb.close()
    return out.getvalue()


def import_rows(content, filename):
    if len(content) > 2*1024*1024:
        raise ValueError("Fichier limité à 2 Mo.")
    if filename.lower().endswith(".csv"):
        frame = pd.read_csv(io.BytesIO(content), dtype={"identifiant":str}, nrows=101)
    elif filename.lower().endswith(".xlsx"):
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        name = "Fouilles" if "Fouilles" in wb.sheetnames else wb.sheetnames[0]
        wb.close()
        frame = pd.read_excel(io.BytesIO(content), sheet_name=name, dtype={"identifiant":str}, nrows=101)
    else:
        raise ValueError("Utiliser un CSV ou un classeur .xlsx.")
    if len(frame) > 100 or len(frame) == 0:
        raise ValueError("Importer entre 1 et 100 fouilles.")
    if not set(ROW_KEYS).issubset(frame.columns):
        raise ValueError("Colonnes attendues : "+", ".join(ROW_KEYS))
    return validate_rows(frame[list(ROW_KEYS)].astype(object).to_dict("records"))


def import_project(content):
    if len(content) > 2*1024*1024:
        raise ValueError("Fichier limité à 2 Mo.")
    def reject_constant(value):
        raise ValueError("Valeur JSON non finie : "+value)
    return validate_project(json.loads(content.decode("utf-8-sig"), parse_constant=reject_constant))


def result_sheets(c):
    sheets = {"Projet":pd.DataFrame({"parametre":["Titre","Origine fouille","Origine lot","Capacité lot (m³)"],
        "valeur":[c["titre"],c["origine"],c["origine_fouilles"],str(c["capacite_lot_m3"])]})}
    sheets["Fouilles"] = pd.DataFrame(c["fouilles"], columns=ROW_KEYS)
    if c["fouille"] is not None:
        sheets["Fouille_principale"] = pd.DataFrame([c["fouille"] | calculate(c["fouille"],c["capacite_m3"])])
    if c["fouilles"]:
        batch = batch_results(c["fouilles"], c["capacite_lot_m3"])
        sheets["Resultats_fouilles"] = batch["table"]
        sheets["Transport_lot"] = pd.DataFrame([{"volume_m3":batch["foisonne_m3"],"voyages_regroupes":batch["voyages"],"voyages_separes":batch["separes"]}])
    if c["comparaison_active"]:
        sheets["Comparaison"] = pd.DataFrame(compare(c), index=["Actuel","Variante"]).reset_index(names="scenario")
    sheets["Parametres_simulation"] = pd.DataFrame([c["incertitude"] | {"executee":c["simulation_active"]}])
    if c["simulation_active"]:
        sim = cached_simulation(payload(c))
        sheets["Simulation"] = pd.DataFrame([{k:sim[k] for k in ("probabilite","analytique","ic_bas","ic_haut","moyenne_volume","moyenne_voyages")}])
    return sheets


def diagram_svg(pit):
    p = pit or default_pit()
    value = lambda key: f"{p[key]:g} m" if pit is not None else {"longueur_m":"L","largeur_m":"B","profondeur_m":"h"}[key]
    # Plan preserves L/B; the section uses independent horizontal and vertical scales.
    scale = min(525/p["longueur_m"], 100/p["largeur_m"])
    w, h = p["longueur_m"]*scale, p["largeur_m"]*scale
    x, y = 370-w/2, 334-h/2
    if w < 3 or h < 3:
        plan = '<text x="80" y="340" fill="#56616B" font-size="14">Rapport L/B extrême : plan réduit à une ligne.</text>'
    else:
        plan = f'<rect x="{x}" y="{y}" width="{w}" height="{h}" fill="#E5F0E8" stroke="{GREEN}" stroke-width="2"/><line x1="{x}" x2="{x+w}" y1="{y+h+18}" y2="{y+h+18}" stroke="{GREEN}"/><text x="370" y="{y+h+38}" text-anchor="middle" font-size="14" fill="{GREEN}">L = {value("longueur_m")}</text><text x="{x+w+12}" y="{y+h/2}" font-size="14" fill="{GREEN}">B = {value("largeur_m")}</text>'
    depth = 96 if p["profondeur_m"] else 0
    bottom = 96+depth
    return f'''<div class="drawing" style="overflow-x:auto;background:#fff;border:1px solid #DFE6E1;border-radius:6px;">
    <svg role="img" aria-label="Coupe et plan de la fouille rectangulaire" viewBox="0 0 760 465" xmlns="http://www.w3.org/2000/svg" style="display:block;width:100%;min-width:510px;height:auto;font-family:Arial,sans-serif;">
    <defs><pattern id="earth" width="14" height="14" patternUnits="userSpaceOnUse"><path d="M0 14L14 0" stroke="#D3E4D7" stroke-width="1"/></pattern><marker id="arrow" viewBox="0 0 10 10" refX="5" refY="5" markerWidth="5" markerHeight="5" orient="auto-start-reverse"><path d="M0 0L10 5L0 10" fill="{BLUE}"/></marker></defs>
    <rect width="760" height="465" fill="#fff"/>
    <text x="25" y="31" font-size="13" font-weight="bold" fill="{GREEN}">01 / COUPE LONGITUDINALE</text>
    <text x="735" y="31" text-anchor="end" font-size="12" fill="#56616B">Axes à échelles indépendantes</text>
    <path d="M55 96H185V{bottom}H575V96H705V230H55Z" fill="#EEF5F0"/>
    <path d="M55 96H185V{bottom}H575V96H705V230H55Z" fill="url(#earth)"/>
    <path d="M55 96H185V{bottom}H575V96H705" fill="none" stroke="{GREEN}" stroke-width="3"/>
    <line x1="185" x2="575" y1="96" y2="96" stroke="#9CB5A5" stroke-dasharray="5 4"/>
    <line x1="185" x2="575" y1="68" y2="68" stroke="{BLUE}" marker-start="url(#arrow)" marker-end="url(#arrow)"/>
    <text x="380" y="58" text-anchor="middle" fill="{BLUE}" font-size="15" font-weight="bold">L = {value("longueur_m")}</text>
    <line x1="615" x2="615" y1="100" y2="{max(bottom-4,101)}" stroke="{BLUE}" marker-start="url(#arrow)" marker-end="url(#arrow)"/>
    <text x="630" y="{115+depth/2}" fill="{BLUE}" font-size="14">h = {value("profondeur_m")}</text>
    <text x="380" y="{min(bottom-12,175)}" text-anchor="middle" fill="#56616B" font-size="14">{'Volume excavé' if depth else 'Profondeur nulle'}</text>
    <text x="25" y="271" font-size="13" font-weight="bold" fill="{GREEN}">02 / VUE EN PLAN</text>
    <text x="735" y="271" text-anchor="end" font-size="12" fill="#56616B">Proportions L/B conservées</text>{plan}
    <text x="25" y="447" fill="#56616B" font-size="12">Parois verticales : géométrie pédagogique, stabilité non étudiée.</text></svg></div>'''


DRAWING = st.components.v2.component("terrebalance_drawing", html='<div id="root"></div>',
    css=':host{display:block;max-width:100%} #root{min-width:0;max-width:100%}',
    js='export default function({data,parentElement}){parentElement.querySelector("#root").innerHTML=data.html;}')


CSS = '''<style>
  .stApp {background:#F5F7FA;color:#202427}
  [data-testid="stMainBlockContainer"] {max-width:1460px;padding:4.3rem 2.5rem 2rem}
  h1,h2,h3,p,span {letter-spacing:0}
  h1 {font-size:2rem!important} h2 {font-size:1.45rem!important} h3 {font-size:1.13rem!important}
  .tb-brand {display:flex;gap:15px;align-items:center;border-bottom:1px solid #CDD8D0;padding-bottom:22px;margin-bottom:6px;flex-wrap:wrap}
  .tb-mark {font-size:19px;font-weight:800;color:#fff;background:#166534;border-radius:5px;padding:13px 10px}
  .tb-brand h1 {margin:0;padding:0;line-height:1.3;color:#143921}
  .tb-brand p {margin:3px 0 0;color:#56616B;font-size:14px}
  .tb-status {margin-left:auto;max-width:100%;font-size:12px;border-left:3px solid #D94841;padding:6px 12px;color:#71312D;overflow-wrap:anywhere}
  .tb-results {display:grid;grid-template-columns:repeat(3,minmax(0,1fr));gap:12px;margin:14px 0 20px}
  .tb-result {padding:12px 14px;background:#fff;border-left:4px solid var(--accent);border-radius:4px;min-width:0}
  .tb-result strong {display:block;font-size:24px;color:var(--accent);overflow-wrap:anywhere}
  .tb-result span {font-size:13px;color:#56616B}
  [data-testid="stForm"] {border:1px solid #DCE3DF;border-radius:6px;background:white}
  .st-key-controls_single,.st-key-controls_variant,.st-key-controls_uncertainty {background:white;border-radius:6px}
  .stButton > button,.stDownloadButton > button,[data-testid="stFormSubmitButton"] button {border-radius:5px!important;min-height:40px}
  button p {white-space:normal}
  button[kind="primary"], [data-testid="stBaseButton-primary"], [data-testid="stBaseButton-primaryFormSubmit"] {background:#166534!important;border-color:#166534!important;color:#fff!important}
  button[kind="primary"] *, [data-testid="stBaseButton-primary"] *,[data-testid="stBaseButton-primaryFormSubmit"] * {color:#fff!important}
  [role="tablist"] {gap:6px;border-bottom:1px solid #CDD8D0;margin-bottom:16px}
  [role="tab"] {padding:9px 15px;min-height:44px;color:#56616B;border-radius:4px 4px 0 0}
  [role="tab"][aria-selected="true"] {color:#166534!important;background:#E5F0E8;font-weight:700}
  [data-baseweb="tab-highlight"] {background:#166534}
  [data-testid="stCaptionContainer"] p {color:#56616B!important;opacity:1}
  .tb-empty {padding:14px 0;color:#56616B;border-bottom:1px solid #DFE5E1}
  @media(max-width:800px) {[data-testid="stMainBlockContainer"] {padding:4.3rem 1.1rem 1.5rem}.tb-status{margin-left:0}}
  @media(max-width:640px) {
    [data-testid="stColumn"] {width:100%;min-width:0;flex:1 1 100%}
    .tb-results {grid-template-columns:1fr;gap:7px}.tb-result strong {font-size:22px}
    .tb-brand {gap:10px;padding-bottom:15px}.tb-brand h1 {font-size:1.8rem!important}
    [role="tablist"] {flex-wrap:wrap;height:auto;overflow:visible;gap:0}
    [role="tab"] {padding:8px 9px;font-size:13px}
    input {font-size:16px!important}
  }
</style>'''


def fmt(value, digits=2):
    return f"{value:,.{digits}f}".replace(","," ").replace(".",",")


def metrics(items):
    content = ''.join(f'<div class="tb-result" style="--accent:{color}"><span>{escape(title)}</span><strong>{escape(value)}</strong></div>' for title,value,color in items)
    st.html('<div class="tb-results">'+content+'</div>')


def chart(fig, key, height=310):
    fig.update_layout(template="plotly_white", height=height, paper_bgcolor="rgba(0,0,0,0)",
        font=dict(family="Arial",size=12,color=INK), colorway=[GREEN,BLUE,CORAL],
        margin=dict(l=15,r=15,t=30,b=65), legend=dict(orientation="h",y=-.2))
    fig.update_xaxes(automargin=True)
    fig.update_yaxes(automargin=True)
    st.plotly_chart(fig, key=key, width="stretch", config={"displayModeBar":False})


def volumes_chart(result):
    return go.Figure(go.Bar(x=["En place","Foisonné"], y=[result["en_place_m3"],result["foisonne_m3"]],
        marker_color=[GREEN,BLUE], text=[fmt(result["en_place_m3"]),fmt(result["foisonne_m3"])], textposition="outside",
        hovertemplate="%{x} : %{y:.4f} m³<extra></extra>")).update_layout(yaxis_title="Volume (m³)",showlegend=False)


def transport_chart(result):
    n, cap = result["voyages"], result["capacite_m3"]
    if n <= 18:
        loads = [cap]*(max(n-1,0))+([result["dernier_m3"]] if n else [])
        fig = go.Figure(go.Bar(x=[str(i+1) for i in range(n)],y=loads,marker_color=CORAL,name="Chargé",
            hovertemplate="Voyage %{x}<br>%{y:.4f} m³<extra></extra>"))
        if n:
            fig.add_hline(y=cap,line_dash="dot",line_color=MUTED)
        fig.update_layout(xaxis_title="Voyage",yaxis_title="Chargement (m³)")
    else:
        last_partial = result["dernier_m3"] if result["partiel"] else 0.
        fig = go.Figure(go.Bar(x=["Chargements complets\n(regroupés)","Dernier partiel"],
            y=[result["complets"]*cap,last_partial],marker_color=[CORAL,"#F1A09A"],
            text=[str(result["complets"])+" voyages", "1 voyage" if last_partial else "Aucun"],textposition="auto"))
        fig.update_layout(yaxis_title="Volume cumulé (m³)")
    return fig


def commit(candidate, message):
    st.session_state.tb_project = validate_project(candidate)
    st.session_state.tb_revision = st.session_state.get("tb_revision",0)+1
    st.session_state.tb_flash = message


def mark_changed(origin):
    return "Exemple fictif modifié" if origin.startswith("Exemple fictif") else "Saisie utilisateur"


def load_example():
    commit(example_project(), "Scénario fictif chargé : fouille de 50 m³ et lot de deux fouilles.")


def reset():
    commit(new_project(), "Nouveau projet vide.")


def pit_inputs(p, prefix):
    cols = st.columns(2)
    with cols[0]:
        length = st.number_input("Longueur L (m)",.01,1000.,float(p["longueur_m"]),.5,key=prefix+"L")
        depth = st.number_input("Profondeur h (m)",0.,100.,float(p["profondeur_m"]),.1,key=prefix+"h")
    with cols[1]:
        width = st.number_input("Largeur B (m)",.01,1000.,float(p["largeur_m"]),.5,key=prefix+"B")
        swell = st.number_input("Foisonnement k (×)",1.,3.,float(p["foisonnement"]),.05,format="%.3f",key=prefix+"k",help="Coefficient multiplicateur. 1,20 = volume augmenté de 20 %.")
    return dict(longueur_m=length,largeur_m=width,profondeur_m=depth,foisonnement=swell)


def earthwork_view(c, rev):
    st.subheader("Une fouille, du terrain à la benne")
    left, right = st.columns([.9,1.5],gap="large")
    with left:
        with st.container(border=True,key='controls_single'):
            pit = pit_inputs(c["fouille"] or default_pit(),f"main_{rev}_")
            cap = st.number_input("Capacité utile C (m³)",.1,1000.,float(c["capacite_m3"]),.5,key=f"main_cap_{rev}")
            dirty=pit != c['fouille'] or cap != c['capacite_m3']
            if dirty:st.badge('Saisie non appliquée',icon=':material/edit:',color='orange')
            st.caption("Bornes de saisie de l’outil, non normatives.")
            if st.button("Calculer la fouille",icon=":material/calculate:",type="primary",width="stretch"):
                draft = copy.deepcopy(c)
                draft.update(fouille=pit,capacite_m3=cap,origine=mark_changed(c["origine"]))
                if dirty:draft.update(comparaison_active=False,simulation_active=False)
                commit(draft,"Fouille enregistrée."+(' Comparaison et simulation à relancer.' if dirty else ''))
                st.rerun()
        with st.expander("Formules et hypothèses",icon=":material/functions:"):
            st.markdown("**V = L × B × h** · **Vf = V × k** · **N = plafond(Vf / C)**")
            st.markdown("Le foisonnement augmente le volume après excavation : **k = 1,20 signifie +20 %**. Aucune valeur n’est attribuée automatiquement à un sol.")
            st.caption("Les résultats et rapports concernent uniquement les valeurs enregistrées. Des champs modifiés dans le formulaire ne sont pas encore appliqués.")
    with right:
        if c["fouille"] is None:
            st.html('<div class="tb-empty">Aucune fouille enregistrée.</div>')
        else:
            result = calculate(c["fouille"],c["capacite_m3"])
            metrics([("Volume en place",fmt(result["en_place_m3"])+" m³",GREEN),
                ("Volume foisonné",fmt(result["foisonne_m3"])+" m³",BLUE),("Transport",str(result["voyages"])+" voyages",CORAL)])
        DRAWING(data={"html":diagram_svg(c["fouille"])},key="pit_drawing",height="content",width="stretch")
    if c["fouille"] is not None:
        result = calculate(c["fouille"],c["capacite_m3"])
        col1,col2 = st.columns(2,gap="large")
        with col1:
            st.subheader("Changement de volume")
            chart(volumes_chart(result),"single_volumes")
            st.caption(f"Augmentation : {fmt(result['augmentation_m3'])} m³, soit {fmt(result['augmentation_pct'],1)} %.")
        with col2:
            st.subheader("Répartition des chargements")
            chart(transport_chart(result),"single_transport")
            st.caption(f"{result['complets']} chargement(s) complet(s). Dernier chargement : {fmt(result['dernier_m3'])} m³ ({fmt(result['dernier_pct'],1)} %)." if result["voyages"] else "Aucun chargement : profondeur nulle.")


def batch_view(c, rev):
    st.subheader("Un lot de fouilles")
    st.caption(c["origine_fouilles"]+" · volumes sans mutualisation de sols ni contrainte de masse")
    with st.form(f"batch_{rev}"):
        edited = st.data_editor(pd.DataFrame(c["fouilles"],columns=ROW_KEYS),num_rows="dynamic",hide_index=True,width="stretch",key=f"batch_editor_{rev}",
            column_config={"identifiant":st.column_config.TextColumn("Identifiant",max_chars=40)} | {
                k:st.column_config.NumberColumn(title,min_value=lo,max_value=hi,format="%.3f") for k,title,lo,hi in [
                    ("longueur_m","L (m)",.01,1000.),("largeur_m","B (m)",.01,1000.),("profondeur_m","h (m)",0.,100.),("foisonnement","k (×)",1.,3.)]})
        cap = st.number_input("Capacité de la benne du lot (m³)",.1,1000.,float(c["capacite_lot_m3"]),.5)
        if st.form_submit_button("Enregistrer les fouilles",icon=":material/save:",type="primary"):
            draft = copy.deepcopy(c)
            draft.update(fouilles=edited.astype(object).to_dict("records"),capacite_lot_m3=cap,origine_fouilles=mark_changed(c["origine_fouilles"]))
            commit(draft,"Lot enregistré.")
            st.rerun()
    if not c["fouilles"]:
        st.info("Aucune fouille dans le lot.")
        return
    result = batch_results(c["fouilles"],c["capacite_lot_m3"])
    metrics([("Volume foisonné total",fmt(result["foisonne_m3"])+" m³",BLUE),
        ("Transport regroupé",str(result["voyages"])+" voyages",GREEN),("Transport séparé",str(result["separes"])+" voyages",CORAL)])
    st.caption("Regroupé : chargements partagés entre fouilles. Séparé : aucun chargement partagé ; arrondi au supérieur pour chaque fouille. Le regroupement n’est qu’une hypothèse logistique.")
    table = result["table"]
    fig = go.Figure(go.Bar(x=table.identifiant,y=table.foisonne_m3,marker_color=BLUE,hovertemplate="%{x} : %{y:.4f} m³<extra></extra>"))
    fig.update_layout(xaxis_title="Fouille",yaxis_title="Volume foisonné (m³)")
    chart(fig,"batch_contribution")
    st.dataframe(table[list(ROW_KEYS)+["en_place_m3","foisonne_m3","voyages"]],hide_index=True,width="stretch")
    with st.expander("Pourquoi le nombre de voyages change-t-il ?"):
        st.markdown("Deux volumes de **12 m³** et une benne de **10 m³** : séparément, 2 + 2 = **4 voyages** ; regroupés, plafond(24/10) = **3 voyages**.")


def comparison_view(c, rev):
    st.subheader("Deux scénarios, les mêmes dimensions en plan")
    if c["fouille"] is None:
        st.info("Aucune fouille principale enregistrée.")
        return
    st.caption(f"Longueur {fmt(c['fouille']['longueur_m'])} m · largeur {fmt(c['fouille']['largeur_m'])} m · scénario actuel : {c['origine']}")
    with st.container(border=True,key='controls_variant'):
        v = c["variante"]
        cols = st.columns(3)
        with cols[0]:h=st.number_input("Profondeur variante (m)",0.,100.,float(v["profondeur_m"]),.1)
        with cols[1]:k=st.number_input("Foisonnement variante (×)",1.,3.,float(v["foisonnement"]),.05)
        with cols[2]:cap=st.number_input("Capacité variante (m³)",.1,1000.,float(v["capacite_m3"]),.5)
        if not c['comparaison_active'] or dict(profondeur_m=h,foisonnement=k,capacite_m3=cap) != c['variante']:
            st.badge('Variante non appliquée',icon=':material/edit:',color='orange')
        if st.button("Comparer les scénarios",icon=":material/compare_arrows:",type="primary"):
            draft=copy.deepcopy(c)
            draft.update(variante=dict(profondeur_m=h,foisonnement=k,capacite_m3=cap),comparaison_active=True)
            commit(draft,"Variante enregistrée et comparée.")
            st.rerun()
    if not c["comparaison_active"]:
        st.caption("Aucune comparaison enregistrée pour la fouille actuelle.")
        return
    current, variant = compare(c)
    keys=["en_place_m3","foisonne_m3","augmentation_m3","capacite_m3","voyages","dernier_m3","dernier_pct"]
    labels=['Volume en place (m³)','Volume foisonné (m³)','Augmentation (m³)','Capacité (m³)','Voyages','Dernier chargement (m³)','Dernier remplissage (%)']
    st.dataframe(pd.DataFrame({"Grandeur":labels,"Actuel":[current[k] for k in keys],"Variante":[variant[k] for k in keys],
        "Écart variante − actuel":[variant[k]-current[k] for k in keys]}),hide_index=True,width="stretch")
    fig=go.Figure()
    for name,key,color in [('En place','en_place_m3',GREEN),('Foisonné','foisonne_m3',BLUE)]:
        fig.add_trace(go.Bar(x=['Actuel','Variante'],y=[current[key],variant[key]],name=name,marker_color=color))
    fig.update_layout(barmode="group",yaxis_title="Volume (m³)")
    chart(fig,"comparison_volumes")
    st.warning("Moins de voyages ne signifie pas une solution techniquement préférable : charge en masse, accès et stabilité ne sont pas vérifiés.")


def uncertainty_view(c, rev):
    st.subheader("La profondeur varie : que devient le transport ?")
    if c["fouille"] is None:
        st.info("Aucune fouille principale enregistrée.")
        return
    with st.container(border=True,key='controls_uncertainty'):
        u=c["incertitude"]
        cols=st.columns(3)
        with cols[0]:a=st.number_input("Profondeur minimale (m)",0.,100.,float(u["min_m"]),.05)
        with cols[1]:m=st.number_input("Profondeur la plus probable (m)",0.,100.,float(u["mode_m"]),.05)
        with cols[2]:b=st.number_input("Profondeur maximale (m)",0.,100.,float(u["max_m"]),.05)
        cols=st.columns(3)
        with cols[0]:n=st.number_input("Nombre de tirages",100,50000,int(u["tirages"]),1000)
        with cols[1]:seed=st.number_input("Graine aléatoire",0,2147483647,int(u["graine"]))
        with cols[2]:limit=st.number_input("Nombre de voyages à ne pas dépasser",0,3000000000,int(u["limite_voyages"]))
        pending=dict(min_m=a,mode_m=m,max_m=b,tirages=int(n),graine=int(seed),limite_voyages=int(limit))
        if not c['simulation_active'] or pending != c['incertitude']:
            st.badge('Hypothèses non appliquées',icon=':material/edit:',color='orange')
        if st.button("Simuler le transport",icon=":material/play_arrow:",type="primary"):
            draft=copy.deepcopy(c)
            draft.update(incertitude=pending,simulation_active=True)
            commit(draft,"Simulation enregistrée. Les tirages ne sont pas des observations de terrain.")
            st.rerun()
    st.caption("Hypothèse triangulaire sur h uniquement. Longueur, largeur, foisonnement et capacité sont fixes. Le mode de la loi n’est pas nécessairement la profondeur nominale.")
    if not c["simulation_active"]:
        st.caption("Aucune simulation exécutée pour la fouille actuelle.")
        return
    sim=cached_simulation(payload(c))
    metrics([("Risque simulé N > "+str(c["incertitude"]["limite_voyages"]),fmt(100*sim["probabilite"],2)+" %",CORAL),
        ("Probabilité analytique",fmt(100*sim["analytique"],2)+" %",GREEN),("Moyenne des voyages",fmt(sim["moyenne_voyages"],3),BLUE)])
    if sim["constant"]:
        st.caption("Profondeur constante : cas déterministe, sans incertitude Monte-Carlo.")
    else:
        st.caption(f"IC binomial exact à 95 % du risque estimé : [{fmt(100*sim['ic_bas'])} % ; {fmt(100*sim['ic_haut'])} %]. Ce n’est pas un intervalle de volumes.")
    df=sim["table"]
    cols=st.columns(2,gap="large")
    with cols[0]:
        fig=go.Figure(go.Histogram(x=df.volume_foisonne_m3,nbinsx=35,marker_color=BLUE))
        fig.update_layout(xaxis_title="Volume foisonné simulé (m³)",yaxis_title="Tirages")
        chart(fig,"simulation_volume")
    with cols[1]:
        counts=df.voyages.value_counts().sort_index()
        if len(counts)<=30:
            fig=go.Figure(go.Bar(x=[str(x) for x in counts.index],y=counts.values/len(df),marker_color=CORAL))
            fig.update_layout(xaxis_title="Nombre entier de voyages")
        else:
            edges=np.linspace(int(df.voyages.min())-.5,int(df.voyages.max())+.5,21)
            values,edges=np.histogram(df.voyages,bins=edges)
            fig=go.Figure(go.Bar(x=[f"{math.ceil(edges[i])}–{math.ceil(edges[i+1])-1}" for i in range(20)],y=values/len(df),marker_color=CORAL))
            fig.update_layout(xaxis_title="Classes de nombres de voyages")
        fig.update_layout(yaxis_title="Fréquence",yaxis_tickformat=".0%")
        chart(fig,"simulation_trips")
    nominal=calculate(c["fouille"],c["capacite_m3"])
    st.dataframe(pd.DataFrame({"Indicateur":["Volume nominal (m³)","Volume moyen simulé (m³)","Volume moyen théorique (m³)","Voyages nominaux","Voyages moyens simulés"],
        "Valeur":[nominal["foisonne_m3"],sim["moyenne_volume"],sim["volume_moyen_theorique"],nominal["voyages"],sim["moyenne_voyages"]]}),hide_index=True,width="stretch")
    st.dataframe(pd.DataFrame({"Quantile":["P5","P50","P95"],"Volume simulé (m³)":sim["quantiles"],"Volume théorique (m³)":sim["quantiles_theoriques"]}),hide_index=True,width="stretch")
    st.download_button("Tirages CSV",safe_csv(df),file_name="terrebalance-tirages.csv",mime="text/csv",icon=":material/download:",on_click='ignore')
    with st.expander("Interpréter la simulation",icon=":material/functions:"):
        st.markdown("La loi triangulaire traduit une hypothèse minimum/mode/maximum, pas une propriété universelle du sol. Les quantiles décrivent une distribution, sans garantir une quantité maximale.")
        st.markdown("**Moyenne des plafonds ≠ plafond de la moyenne.** La moyenne des voyages peut être non entière ; chaque scénario simulé comporte un nombre entier de voyages.")
        st.caption("Les seuils proches d’un entier sont stabilisés sur huit unités de précision flottante. Le contrôle N > n est cohérent avec Vf > n × C selon cette même convention, sans arrondir les volumes affichés pour décider.")


def reference_tests():
    tests=[]
    def check(name, expected, actual, success):
        tests.append(dict(test=name,attendu=expected,obtenu=actual,reussi=bool(success)))
    r=calculate(default_pit(),10)
    check("Volume et voyages", "50 / 60 m³ / 6 voyages",f"{r['en_place_m3']:g} / {r['foisonne_m3']:g} m³ / {r['voyages']} voyages",r['en_place_m3']==50 and r['foisonne_m3']==60 and r['voyages']==6)
    r=calculate(default_pit(),8)
    check("Dernier chargement", "8 voyages / 4 m³ / 50 %",f"{r['voyages']} / {r['dernier_m3']:g} m³ / {r['dernier_pct']:g} %",r['voyages']==8 and r['dernier_m3']==4 and r['dernier_pct']==50)
    r=calculate(default_pit()|dict(profondeur_m=0),10)
    check("Profondeur nulle", "0 m³ / 0 voyage",f"{r['foisonne_m3']:g} m³ / {r['voyages']} voyage",r['voyages']==0 and r['dernier_m3']==0)
    r=batch_results(example_project()['fouilles'],10)
    check("Regroupement", "3 regroupés / 4 séparés",f"{r['voyages']} / {r['separes']}",r['voyages']==3 and r['separes']==4)
    c=example_project()
    s=simulate(c)
    check("Loi triangulaire", "P(N > 6) = 50 %",f"{100*s['analytique']:.8f} %",abs(s['analytique']-.5)<1e-12)
    c['incertitude'].update(min_m=1.,mode_m=1.,max_m=1.)
    s=simulate(c)
    check("Profondeur constante", "6 voyages / risque nul",f"{s['moyenne_voyages']:g} / {s['probabilite']:g}",s['probabilite']==0 and s['moyenne_voyages']==6)
    return tests


def pdf_report(configuration):
    c=validate_project(json.loads(configuration))
    font_dir=Path(matplotlib.get_data_path())/'fonts'/'ttf'
    for name,file in [('TB','DejaVuSans.ttf'),('TB-Bold','DejaVuSans-Bold.ttf')]:
        if name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(name,str(font_dir/file)))
    pdfmetrics.registerFontFamily('TB',normal='TB',bold='TB-Bold',italic='TB',boldItalic='TB-Bold')
    styles=getSampleStyleSheet()
    body=ParagraphStyle('TBbody',parent=styles['Normal'],fontName='TB',fontSize=9,leading=13,spaceAfter=7,textColor=colors.HexColor(INK))
    heading=ParagraphStyle('TBheading',parent=body,fontName='TB-Bold',fontSize=12,leading=16,spaceBefore=12,spaceAfter=7,keepWithNext=True,textColor=colors.HexColor(GREEN))
    title=ParagraphStyle('TBtitle',parent=heading,fontSize=25,leading=30)
    cell=ParagraphStyle('TBcell',parent=body,fontSize=7.3,leading=10,spaceAfter=0)
    cellhead=ParagraphStyle('TBcellhead',parent=cell,fontName='TB-Bold',textColor=colors.white)
    story=[Paragraph('TerreBalance',title),Paragraph('Volumes de terrassement et transport des terres',heading),
        Paragraph(escape(c['titre']),body),Paragraph('Édité le '+datetime.now().astimezone().strftime('%d/%m/%Y %H:%M %z'),body),
        Paragraph(escape(LIMITS),body)]
    def para(text):
        story.append(Paragraph(text,body))
    def table(headers,rows,widths):
        data=[[Paragraph(escape(str(v)),cellhead if index==0 else cell) for v in row] for index,row in enumerate([headers]+rows)]
        t=Table(data,colWidths=[w*mm for w in widths],repeatRows=1,hAlign='LEFT')
        t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor(GREEN)),
            ('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor(BG)]),
            ('GRID',(0,0),(-1,-1),.4,colors.HexColor('#DCE3DF')),('VALIGN',(0,0),(-1,-1),'TOP'),
            ('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5)]))
        story.extend([t,Spacer(1,7)])
    def picture(fig,name):
        try:
            buffer=io.BytesIO()
            fig.savefig(buffer,format='png',dpi=150,bbox_inches='tight')
            buffer.seek(0)
            im=Image(buffer)
            aspect=im.imageHeight/im.imageWidth
            story.append(KeepTogether([Paragraph(name,heading),Image(buffer,width=174*mm,height=174*mm*aspect),Spacer(1,7)]))
        finally:
            plt.close(fig)
    if c['fouille'] is not None:
        p=c['fouille'];r=calculate(p,c['capacite_m3'])
        story.append(Paragraph('1. Fouille principale',heading))
        para('Origine : '+escape(c['origine']))
        table(['L (m)','B (m)','h (m)','k (×)','C (m³)'],[[fmt(p[k],3) for k in PIT_KEYS]+[fmt(c['capacite_m3'])]],[35,35,35,35,35])
        table(['Résultat','Valeur'],[
            ['Volume en place',fmt(r['en_place_m3'])+' m³'],['Volume foisonné',fmt(r['foisonne_m3'])+' m³'],
            ['Augmentation',fmt(r['augmentation_m3'])+' m³ / '+fmt(r['augmentation_pct'],1)+' %'],
            ['Voyages',r['voyages']],['Chargements complets',r['complets']],
            ['Dernier chargement',fmt(r['dernier_m3'])+' m³ / '+fmt(r['dernier_pct'],1)+' %']],[88,87])
        fig,axes=plt.subplots(1,2,figsize=(9,3.3))
        axes[0].bar(['En place','Foisonné'],[r['en_place_m3'],r['foisonne_m3']],color=[GREEN,BLUE])
        axes[0].set_ylabel('Volume (m³)')
        if r['voyages']<=18:
            loads=[c['capacite_m3']]*max(r['voyages']-1,0)+([r['dernier_m3']] if r['voyages'] else [])
            axes[1].bar(range(1,len(loads)+1),loads,color=CORAL)
            axes[1].set_xlabel('Voyage')
        else:
            axes[1].bar(['Complets\n(regroupés)','Dernier partiel'],[r['complets']*c['capacite_m3'],r['dernier_m3'] if r['partiel'] else 0],color=CORAL)
        axes[1].set_ylabel('Chargement (m³)')
        for ax in axes:ax.spines[['top','right']].set_visible(False)
        fig.tight_layout()
        picture(fig,'Volumes et chargements')
    if c['fouilles']:
        story.append(Paragraph('2. Lot de fouilles',heading))
        para('Origine : '+escape(c['origine_fouilles'])+' ; capacité commune : '+fmt(c['capacite_lot_m3'])+' m³.')
        result=batch_results(c['fouilles'],c['capacite_lot_m3'])
        rows=[]
        for row in result['table'].to_dict('records'):
            rows.append([row['identifiant']]+[fmt(row[k],3) for k in PIT_KEYS]+[fmt(row['foisonne_m3']),row['voyages']])
        table(['Fouille','L (m)','B (m)','h (m)','k (×)','Vf (m³)','Voyages séparés'],rows,[37,21,21,21,21,28,26])
        para(f"Volume foisonné total : <b>{fmt(result['foisonne_m3'])} m³</b>. Transport regroupé : <b>{result['voyages']} voyages</b>. Transport séparé : <b>{result['separes']} voyages</b>.")
        para('Le regroupement suppose des chargements partagés entre fouilles. Il ne vérifie pas la possibilité réelle de mélanger les matériaux ou les contraintes logistiques.')
    if c['comparaison_active']:
        story.append(Paragraph('3. Comparaison enregistrée',heading))
        a,b=compare(c)
        table(['Scénario','h (m)','k (×)','C (m³)','Vf (m³)','Voyages'],[
            ['Actuel',fmt(c['fouille']['profondeur_m']),fmt(c['fouille']['foisonnement'],3),fmt(c['capacite_m3']),fmt(a['foisonne_m3']),a['voyages']],
            ['Variante',fmt(c['variante']['profondeur_m']),fmt(c['variante']['foisonnement'],3),fmt(c['variante']['capacite_m3']),fmt(b['foisonne_m3']),b['voyages']]],[35,28,28,28,28,28])
        para('Longueur et largeur communes à la fouille principale. Le nombre de voyages ne constitue pas à lui seul un critère de choix technique.')
    if c['simulation_active']:
        story.append(Paragraph('4. Simulation de profondeur',heading))
        s=cached_simulation(configuration);u=c['incertitude']
        para(f"Hypothèse triangulaire : minimum {fmt(u['min_m'],3)} m, mode {fmt(u['mode_m'],3)} m, maximum {fmt(u['max_m'],3)} m. {u['tirages']} tirages, graine {u['graine']}. Les autres paramètres sont fixes.")
        table(['Indicateur','Valeur'],[
            ['P(N > '+str(u['limite_voyages'])+') simulée',fmt(100*s['probabilite'])+' %'],
            ['Probabilité analytique',fmt(100*s['analytique'])+' %'],
            ['IC binomial exact 95 %' if not s['constant'] else 'Cas déterministe',
             '['+fmt(100*s['ic_bas'])+' % ; '+fmt(100*s['ic_haut'])+' %]' if not s['constant'] else 'Profondeur constante'],
            ['Volume moyen simulé',fmt(s['moyenne_volume'],4)+' m³'],['Voyages moyens simulés',fmt(s['moyenne_voyages'],4)],
            ['Quantiles simulés P5 / P50 / P95',' / '.join(fmt(x,4) for x in s['quantiles'])+' m³']],[87,88])
        fig,ax=plt.subplots(figsize=(9,3.1))
        if s['constant']:
            value=float(s['table'].volume_foisonne_m3.iloc[0])
            spread=max(abs(value)*.05,1.)
            ax.axvline(value,color=BLUE,label=f'Volume constant : {value:g} m³')
            ax.set_xlim(value-spread,value+spread)
            ax.set_yticks([])
            ax.legend()
        else:ax.hist(s['table'].volume_foisonne_m3,bins=35,color=BLUE,alpha=.85)
        ax.set_xlabel('Volume foisonné simulé (m³)');ax.set_ylabel('Cas déterministe' if s['constant'] else 'Tirages')
        fig.tight_layout();picture(fig,'Distribution des volumes simulés')
        para('Les tirages sont des résultats simulés, pas des mesures. Les quantiles ne sont pas des garanties ; l’intervalle binomial porte sur la probabilité estimée, pas sur les volumes. Un nombre moyen de voyages peut être non entier.')
    story.append(Paragraph('Méthode et limites',heading))
    para('V = L × B × h ; Vf = V × k ; N = plafond(Vf/C). k = 1,20 représente une augmentation de volume de 20 %. Les limites numériques de saisie sont des choix de l’outil, non des prescriptions normatives.')
    para('Les décisions de seuil utilisent les valeurs non arrondies, avec stabilisation de huit unités de précision flottante autour des rapports entiers. Une profondeur nulle donne zéro volume et zéro voyage.')
    out=io.BytesIO()
    doc=SimpleDocTemplate(out,pagesize=A4,leftMargin=17*mm,rightMargin=17*mm,topMargin=15*mm,bottomMargin=17*mm,title='TerreBalance - Rapport pédagogique')
    def footer(canvas,document):
        canvas.setFont('TB',8);canvas.setFillColor(colors.HexColor(MUTED))
        canvas.drawString(17*mm,9*mm,'TerreBalance | Document pédagogique')
        canvas.drawRightString(A4[0]-17*mm,9*mm,str(document.page))
    doc.build(story,onFirstPage=footer,onLaterPages=footer)
    return out.getvalue()


@st.cache_data(max_entries=3,ttl=1800,show_spinner=False)
def cached_report(configuration):
    """Rapport Unicode, tableaux pagines et figure constante cadree autour du volume."""
    return pdf_report(configuration)


def report_view(c,rev):
    st.subheader("Données et rapport d’étude")
    with st.expander("Nom de l’étude",icon=":material/edit:"):
        with st.form(f'title_{rev}'):
            title=st.text_input('Nom du projet',c['titre'],max_chars=80)
            if st.form_submit_button('Enregistrer le nom',icon=':material/save:'):
                draft=copy.deepcopy(c);draft['titre']=title
                commit(draft,'Nom du projet enregistré.');st.rerun()
    with st.expander("Importer des fouilles ou restaurer un projet",icon=":material/upload:"):
        template=pd.DataFrame(example_project()['fouilles'])
        with st.container(horizontal=True):
            st.download_button('Modèle CSV fictif',safe_csv(template),file_name='terrebalance-modele-fictif.csv',mime='text/csv',icon=':material/download:',on_click='ignore')
            st.download_button('Modèle Excel fictif',excel_bytes({'Fouilles':template}),file_name='terrebalance-modele-fictif.xlsx',mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',icon=':material/download:',on_click='ignore')
        upload=st.file_uploader('Fichier CSV, Excel ou JSON',type=['csv','xlsx','json'],max_upload_size=2)
        if upload:
            try:
                is_json=upload.name.lower().endswith('.json')
                imported=import_project(upload.getvalue()) if is_json else import_rows(upload.getvalue(),upload.name)
                if is_json:
                    st.caption('Projet JSON : '+imported['titre']+' · '+imported['origine'])
                    preview=imported['fouilles'] or ([{'identifiant':'Fouille principale',**imported['fouille']}] if imported['fouille'] else [])
                else:preview=imported
                st.caption('Fichier prêt : '+upload.name)
                st.dataframe(pd.DataFrame(preview),hide_index=True,width='stretch')
                if not is_json:st.caption('Import tabulaire : seul le lot sera remplacé ; la fouille principale et les capacités sont conservées.')
                file_id=hashlib.sha256(upload.getvalue()).hexdigest()
                confirm=st.checkbox('Confirmer le remplacement par le fichier',key='confirm_import_'+file_id)
                if st.button('Appliquer le fichier',disabled=not confirm,icon=':material/upload:'):
                    draft=imported if is_json else copy.deepcopy(c)
                    if not is_json:draft.update(fouilles=imported,origine_fouilles='Import : origine à documenter')
                    commit(draft,'Fichier validé et appliqué.');st.rerun()
            except (ValueError,TypeError,KeyError,UnicodeError,io.UnsupportedOperation) as exc:
                st.error('Import refusé : '+str(exc))
            except Exception:
                st.error('Fichier illisible ou endommagé. Aucun remplacement effectué.')
        st.caption('CSV UTF-8 : virgule séparatrice, point décimal. Maximum 100 fouilles et 2 Mo. Le JSON restaure tous les paramètres, les origines et les états de calcul.')
    st.markdown('#### Projet enregistré')
    with st.container(horizontal=True):
        st.download_button('Projet JSON complet',json.dumps(c,ensure_ascii=False,indent=2,allow_nan=False).encode(),file_name='terrebalance-projet.json',mime='application/json',icon=':material/save:',on_click='ignore')
        st.download_button('Fouilles CSV',safe_csv(pd.DataFrame(c['fouilles'],columns=ROW_KEYS)),file_name='terrebalance-fouilles.csv',mime='text/csv',icon=':material/download:',on_click='ignore')
        st.download_button('Classeur Excel',excel_bytes(result_sheets(c)),file_name='terrebalance-etude.xlsx',mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',icon=':material/download:',on_click='ignore')
    ready=c['fouille'] is not None or bool(c['fouilles'])
    if st.button('Préparer le rapport PDF',disabled=not ready,icon=':material/picture_as_pdf:',type='primary'):
        with st.spinner('Préparation du rapport…'):
            st.session_state.tb_report=(payload(c),cached_report(payload(c)))
    report=st.session_state.get('tb_report')
    if report and report[0]==payload(c):
        st.download_button('Télécharger le rapport PDF',report[1],file_name='terrebalance-rapport.pdf',mime='application/pdf',icon=':material/download:',on_click='ignore')
    st.caption('Données en mémoire de session, sans sauvegarde automatique. Un PDF préparé n’est téléchargeable que pour sa configuration exacte. Les simulations et comparaisons ne sont incluses que si elles sont actives.')
    with st.expander('Méthode et tests de référence',icon=':material/task_alt:'):
        st.markdown('Les volumes sont exprimés en m³, les dimensions en mètres et k est sans unité. Le dernier voyage peut être partiellement chargé ; aucune charge en tonnes n’est déduite de ces volumes.')
        if st.button('Exécuter les six tests',icon=':material/checklist:'):
            st.session_state.tb_checks=reference_tests()
        checks=st.session_state.get('tb_checks')
        if checks:
            st.dataframe(pd.DataFrame(checks),hide_index=True,width='stretch')
            if all(t['reussi'] for t in checks):st.success('6 tests de référence réussis.')
            else:st.error('Un test de référence a échoué.')
    st.caption(LIMITS)


def main():
    st.set_page_config(page_title='TerreBalance | Volumes et transport',page_icon=':material/landscape:',layout='wide',initial_sidebar_state='collapsed')
    st.html(CSS)
    st.session_state.setdefault('tb_project',new_project())
    c=copy.deepcopy(st.session_state.tb_project)
    rev=st.session_state.get('tb_revision',0)
    origin=c['origine'] if c['fouille'] is not None else 'Projet sans fouille enregistrée'
    st.html(f'<header class="tb-brand"><span class="tb-mark">TB</span><div><h1>TerreBalance</h1><p>Volumes et transport des terres</p></div><div class="tb-status">{escape(origin)}<br>{escape(c["titre"])}</div></header>')
    with st.container(horizontal=True):
        confirm=st.checkbox('Autoriser le remplacement du projet',key='allow_replace')
        locked=(c['fouille'] is not None or bool(c['fouilles'])) and not confirm
        st.button('Charger l’exemple fictif',icon=':material/science:',disabled=locked,on_click=load_example)
        st.button('Nouveau projet',icon=':material/note_add:',disabled=locked,on_click=reset)
    if 'tb_flash' in st.session_state:st.toast(st.session_state.pop('tb_flash'),icon=':material/check_circle:')
    tabs=st.tabs(SECTIONS,key='tb_tabs',on_change='rerun')
    views=[earthwork_view,batch_view,comparison_view,uncertainty_view,report_view]
    for tab,view in zip(tabs,views):
        with tab:
            try:view(c,rev)
            except (ValueError,TypeError,KeyError,OverflowError) as exc:
                st.error('Données non appliquées : '+str(exc))


if __name__=='__main__':
    main()
